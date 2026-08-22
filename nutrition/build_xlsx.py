# -*- coding: utf-8 -*-
"""Сборка Excel-таблицы протокола питания БГБКБС."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import plan_data as D

FONT = "Arial"
DARK   = "1F3864"   # заголовки листов
HEAD   = "2E5C8A"   # шапки таблиц
BAND   = "EAF0F7"   # чередование строк
ACCENT = "FFF2CC"   # выделение
RED    = "F8D7DA"
GREEN  = "DDF0E0"
INPUT  = "FFFF00"   # ячейка для заполнения

thin = Side(style="thin", color="B7C6D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_row(ws, row, text, ncols, size=14):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1)
    c.font = Font(name=FONT, size=size, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 26
    return row + 1


def note_row(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1)
    c.font = Font(name=FONT, size=9, italic=True, color="44546A")
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = max(16, 14 * (len(text) // 110 + 1))
    return row + 1


def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 32
    return row + 1


def data_rows(ws, row, rows, widths, band=True, bold_first=False, height_per=14, size=10):
    for r, data in enumerate(rows):
        for i, val in enumerate(data, start=1):
            c = ws.cell(row=row, column=i, value=val)
            c.font = Font(name=FONT, size=size,
                          bold=(bold_first and i == 1))
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            c.border = BORDER
            if band and r % 2 == 1:
                c.fill = PatternFill("solid", fgColor=BAND)
        # оценка высоты строки по самой длинной ячейке
        need = 1
        for i, val in enumerate(data):
            w = widths[i] if i < len(widths) else 20
            txt = str(val)
            lines = sum(max(1, int(len(part) / max(w - 2, 6)) + 1)
                        for part in txt.split("\n"))
            need = max(need, lines)
        ws.row_dimensions[row].height = max(16, need * height_per)
        row += 1
    return row


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ═══════════════════════════════════════════════════════════════ ЛИСТ 1
ws = wb.active
ws.title = "Как читать"
ws.sheet_view.showGridLines = False
W = [22, 62, 46]
set_widths(ws, W)
r = 1
r = title_row(ws, r, D.TITLE, 3, size=15)
r = note_row(ws, r, D.SUBTITLE + ". " + D.SOURCE, 3)
r += 1

r = title_row(ws, r, "ГЛАВНОЕ ПРАВИЛО ДНЯ", 3, size=12)
r = header_row(ws, r, D.PRINCIPLE_HEAD)
r = data_rows(ws, r, D.PRINCIPLE, W, bold_first=True)
r += 1

r = title_row(ws, r, "ЧТО ЕЩЁ ОБЯЗАТЕЛЬНО КАЖДЫЙ ДЕНЬ", 3, size=12)
r = header_row(ws, r, ["Что", "Когда", "Как"])
must = [
    ["КЛЕТЧАТКА", "В каждый приём пищи",
     "Овощи отварные или ошпаренные. Сырыми — только огурец и зелень"],
    ["СОРБЕНТ — псиллиум", "Между приёмами пищи: за 2 ч до еды или через 2 ч после",
     "1 ст. л. на стакан воды, сверху запить ещё стаканом. Отдельно от еды и БАДов"],
    ["ТРАВКИ", "После каждого приёма пищи",
     "Укроп и семя укропа, фенхель, анис, тмин, имбирь, куркума, кинза, солодка, шиповник"],
    ["Масло пшеничных зародышей (витамин Е)", "Сразу после завтрака",
     "1 ч. л. Только с едой, не натощак"],
    ["Витамин А", "С завтраком, из еды",
     "Желток, печень индейки, морковь, тыква, шпинат — обязательно с жиром"],
    ["Омега-3", "Завтрак и обед", "1600 мг с едой, 2 раза в день"],
    ["Цинк 25 мг", "Утром сразу после завтрака", "Или на обед — выбрать один приём"],
    ["Витамин С 1000 мг", "Между приёмами пищи", "С биофлавоноидами"],
    ["ВОДА", "В течение дня", "30 мл на кг веса"],
]
r = data_rows(ws, r, must, W, bold_first=True)
r += 1

# Калькулятор воды — формула
r = title_row(ws, r, "НОРМА ВОДЫ — впиши свой вес в жёлтую ячейку", 3, size=12)
ws.cell(row=r, column=1, value="Вес, кг").font = Font(name=FONT, size=10, bold=True)
ws.cell(row=r, column=1).border = BORDER
wcell = ws.cell(row=r, column=2, value=60)
wcell.font = Font(name=FONT, size=10, bold=True, color="0000FF")
wcell.fill = PatternFill("solid", fgColor=INPUT)
wcell.border = BORDER
wcell.alignment = Alignment(horizontal="left", vertical="center")
ws.cell(row=r, column=3, value="Пример. Впиши свой вес — норма пересчитается сама")\
    .font = Font(name=FONT, size=9, italic=True, color="44546A")
ws.cell(row=r, column=3).border = BORDER
weight_ref = f"$B${r}"
r += 1
ws.cell(row=r, column=1, value="Норма воды, л").font = Font(name=FONT, size=10, bold=True)
ws.cell(row=r, column=1).border = BORDER
f = ws.cell(row=r, column=2, value=f"={weight_ref}*30/1000")
f.font = Font(name=FONT, size=10, bold=True)
f.number_format = "0.0"
f.border = BORDER
ws.cell(row=r, column=3, value="Формула из базы: 30 мл × вес")\
    .font = Font(name=FONT, size=9, italic=True, color="44546A")
ws.cell(row=r, column=3).border = BORDER
r += 2

# Готовая справка, чтобы норма была видна без пересчёта формулы
r = title_row(ws, r, "ГОТОВАЯ СПРАВКА ПО ВОДЕ", 3, size=11)
r = header_row(ws, r, ["Вес, кг", "Норма воды в день", "Ориентир"])
water = [[f"{kg} кг", f"{kg*30/1000:.1f} л".replace(".", ","),
          f"примерно {round(kg*30/1000/0.25)} стаканов по 250 мл"]
         for kg in range(45, 105, 5)]
r = data_rows(ws, r, water, W, bold_first=True)
r = note_row(ws, r, "Считается чистая вода, отдельно от чая, морсов и супа.", 3)
r += 1

r = note_row(ws, r, "ЛЕГЕНДА: жёлтая ячейка — единственная, куда нужно что-то вписать "
                    "(свой вес). Всё остальное в файле — справочное, менять не нужно. "
                    "Листы: «Схема дня» — весь день по часам · «Меню 7 дней» — что "
                    "готовить · «Замены» — чем заменить сахар, глютен и молочку · "
                    "«Покупки» — список по магазинам · «Можно-Нельзя» — короткая "
                    "шпаргалка · «Клетчатка» — как готовить овощи · «Рекомендации» — "
                    "текстом, без таблиц.", 3)
r = note_row(ws, r, D.DISCLAIMER, 3)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 2
ws = wb.create_sheet("Схема дня")
ws.sheet_view.showGridLines = False
W = [14, 52, 64]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "СХЕМА ДНЯ — весь день по часам", 3)
r = note_row(ws, r, "Время ориентировочное: важны не часы, а порядок и промежутки "
                    "в 2 часа вокруг сорбента.", 3)
r += 1
r = header_row(ws, r, D.DAY_HEAD)
start = r
r = data_rows(ws, r, D.DAY_SCHEDULE, W, bold_first=True)
# подсветка приёмов пищи и сорбента
for i, row_data in enumerate(D.DAY_SCHEDULE):
    txt = row_data[1]
    rr = start + i
    fill = None
    if txt.startswith(("ЗАВТРАК", "ОБЕД", "УЖИН")):
        fill = GREEN
    elif "ПСИЛЛИУМ" in txt:
        fill = ACCENT
    if fill:
        for col in range(1, 4):
            ws.cell(row=rr, column=col).fill = PatternFill("solid", fgColor=fill)
r += 1
r = note_row(ws, r, D.HERBS_NOTE, 3)
r += 1
r = title_row(ws, r, "ДОПОЛНИТЕЛЬНО ИЗ БАЗЫ — по согласованию, в основную схему не входит", 3, size=11)
r = header_row(ws, r, D.EXTRA_HEAD)
r = data_rows(ws, r, D.EXTRA, W)
r = note_row(ws, r, D.EXTRA_NOTE, 3)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 3
ws = wb.create_sheet("Меню 7 дней")
ws.sheet_view.showGridLines = False
W = [10, 44, 46, 44]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "МЕНЮ НА 7 ДНЕЙ — без сахара, без глютена, без молочки, без курицы", 4)
r = note_row(ws, r, D.MENU_NOTE, 4)
r += 1
r = header_row(ws, r, D.MENU_HEAD)
r = data_rows(ws, r, D.MENU, W, bold_first=True, height_per=15)
r += 1
r = note_row(ws, r, "ЗАВТРАК — белок и жир, без круп и фруктов. "
                    "ОБЕД — единственный приём с крупой, картофелем или бобовыми. "
                    "УЖИН — только белок и жир. Клетчатка везде отварная или ошпаренная.", 4)
r = note_row(ws, r, "Из базы: мясо примерно дважды в неделю, красное мясо — реже; "
                    "при вздутии всё отварное или паровое и порции меньше.", 4)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 4
ws = wb.create_sheet("Замены")
ws.sheet_view.showGridLines = False
W = [40, 62, 30]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "ЧЕМ ЗАМЕНИТЬ САХАР, ГЛЮТЕН И МОЛОЧКУ", 3)
r += 1
for gname, rows in D.SWAP_GROUPS:
    r = title_row(ws, r, gname, 3, size=12)
    r = header_row(ws, r, D.SWAP_HEAD)
    r = data_rows(ws, r, rows, W)
    r += 1
r = note_row(ws, r, "Гхи — единственное исключение из «без молочки»: по базе сливочное "
                    "масло и гхи разрешены, в гхи нет казеина и лактозы.", 3)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 5
ws = wb.create_sheet("Покупки")
ws.sheet_view.showGridLines = False
W = [18, 58, 13, 14, 15]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "СПИСОК ПОКУПОК — Пятёрочка · Перекрёсток · ВкусВилл", 5)
r += 1
r = header_row(ws, r, D.SHOP_HEAD)
start = r
r = data_rows(ws, r, D.SHOP, W, bold_first=True)
for i, row_data in enumerate(D.SHOP):
    for col in (3, 4, 5):
        c = ws.cell(row=start + i, column=col)
        c.alignment = Alignment(horizontal="center", vertical="center")
        v = str(c.value)
        if v == "есть":
            c.fill = PatternFill("solid", fgColor=GREEN)
        elif v == "нет":
            c.fill = PatternFill("solid", fgColor=RED)
        else:
            c.fill = PatternFill("solid", fgColor=ACCENT)
r += 1
r = note_row(ws, r, D.SHOP_NOTE, 5)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 6
ws = wb.create_sheet("Можно-Нельзя")
ws.sheet_view.showGridLines = False
W = [52, 62]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "ШПАРГАЛКА: НЕЛЬЗЯ / МОЖНО", 2)
r += 1
r = header_row(ws, r, D.YESNO_HEAD)
start = r
r = data_rows(ws, r, D.YESNO, W, band=False)
for i in range(len(D.YESNO)):
    ws.cell(row=start + i, column=1).fill = PatternFill("solid", fgColor=RED)
    ws.cell(row=start + i, column=2).fill = PatternFill("solid", fgColor=GREEN)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 7
ws = wb.create_sheet("Клетчатка")
ws.sheet_view.showGridLines = False
W = [34, 62, 18]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "КЛЕТЧАТКА — отварная или ошпаренная", 3)
r = note_row(ws, r, D.FIBER_NOTE, 3)
r += 1
r = header_row(ws, r, D.FIBER_HEAD)
r = data_rows(ws, r, D.FIBER, W, bold_first=True)

# ═══════════════════════════════════════════════════════════════ ЛИСТ 8
ws = wb.create_sheet("Рекомендации")
ws.sheet_view.showGridLines = False
W = [6, 42, 78]
set_widths(ws, W)
r = 1
r = title_row(ws, r, "ОБЩИЕ РЕКОМЕНДАЦИИ", 3)
r = note_row(ws, r, "Здесь то же самое, но словами — чтобы не заглядывать в таблицы "
                    "каждый день.", 3)
r += 1
r = header_row(ws, r, ["№", "Правило", "Что это значит"])
rows = [[str(i), t, b] for i, (t, b) in enumerate(D.RECOMMENDATIONS, start=1)]
r = data_rows(ws, r, rows, W, bold_first=False)
for rr in range(r - len(rows), r):
    ws.cell(row=rr, column=2).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=rr, column=1).alignment = Alignment(horizontal="center", vertical="top")
r += 1
r = note_row(ws, r, D.DISCLAIMER, 3)

# Печать
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.freeze_panes = None

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "pitanie-bgbkbs-menu.xlsx")
wb.save(out)
print("saved:", out)
